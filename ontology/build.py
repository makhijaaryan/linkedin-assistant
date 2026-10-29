"""Generate the frozen ontology artifacts from ontology.py and schema.py.

Run from the repo root:  python -m ontology.build
Writes ontology/generated/: parsed_cv.schema.json, parsed_job.schema.json,
ontology.json, ONTOLOGY.md, ONTOLOGY.docx (the course's starter-ontology format).
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

from ontology import ontology as O
from ontology.schema import ParsedCV, ParsedJob

OUT = Path("ontology/generated")


def relation_rows():
    for name, sources, targets, desc in O.RELATIONS:
        yield name, " | ".join(sources), " | ".join(targets), desc


def write_json():
    (OUT / "parsed_cv.schema.json").write_text(json.dumps(ParsedCV.model_json_schema(), indent=2))
    (OUT / "parsed_job.schema.json").write_text(json.dumps(ParsedJob.model_json_schema(), indent=2))
    spec = {
        "version": O.VERSION,
        "entities": {n: {"descriptor": d, "attributes": a} for n, (d, a) in O.ENTITIES.items()},
        "relations": [{"name": n, "source": s, "target": t, "descriptor": d} for n, s, t, d in O.RELATIONS],
    }
    (OUT / "ontology.json").write_text(json.dumps(spec, indent=2))


def write_markdown():
    lines = [f"# Job-Market Ontology v{O.VERSION}", "",
             "Generated from `ontology/ontology.py` — edit that file, then run `python -m ontology.build`.", "",
             "## Entity types", "", "| Entity Type | Type Descriptor | Attributes [Type] |", "|---|---|---|"]
    for name, (desc, attrs) in O.ENTITIES.items():
        attr_text = "<br>".join(f"{a} [{t}]" for a, t in attrs.items())
        lines.append(f"| **{name}** | {desc} | {attr_text} |")
    lines += ["", "## Relationships", "",
              "| Relationship | Source Entity → Target Entity | Relationship Descriptor |", "|---|---|---|"]
    for name, src, tgt, desc in relation_rows():
        note = " *(added: Location had no relation in Fig. 2)*" if name in O.ADDED_RELATIONS else ""
        lines.append(f"| **{name}** | {src} → {tgt} | {desc}{note} |")
    (OUT / "ONTOLOGY.md").write_text("\n".join(lines) + "\n")


def write_docx():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"LinkedIn Assistant — Minimal Job-Market Ontology (v{O.VERSION})", level=1)
    doc.add_paragraph(
        "The smallest set of entities, relationships and type descriptors sufficient for reliable "
        "extraction from CVs and job postings, structured representation in a knowledge graph, and "
        "downstream match reasoning. Mirrors the Land/Property Dispute reference: parties become "
        "Candidate and Company, claims become Requirements, evidence becomes source spans.")

    doc.add_heading("Entity Types", level=2)
    t = doc.add_table(rows=1, cols=3, style="Light Grid Accent 1")
    for cell, text in zip(t.rows[0].cells, ["Entity Type", "Type Descriptor", "Attributes [Type]"]):
        cell.text = text
    for name, (desc, attrs) in O.ENTITIES.items():
        row = t.add_row().cells
        row[0].text, row[1].text = name, desc
        row[2].text = "\n".join(f"{a} [{ty}]" for a, ty in attrs.items())

    doc.add_heading("Relationships", level=2)
    t = doc.add_table(rows=1, cols=3, style="Light Grid Accent 1")
    for cell, text in zip(t.rows[0].cells, ["Relationship", "Source Entity → Target Entity", "Relationship Descriptor"]):
        cell.text = text
    for name, src, tgt, desc in relation_rows():
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = name, f"{src} → {tgt}", desc
    doc.save(OUT / "ONTOLOGY.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    write_json()
    write_markdown()
    write_docx()
    print(f"Ontology v{O.VERSION}: {len(O.ENTITIES)} entity types, {len(O.RELATIONS)} relations -> {OUT}/")
